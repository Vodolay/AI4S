"""Render a report Markdown source into Word, in the house table design.

A reviewer's Word export of this report read better than the LaTeX build --
chiefly its tables, which carry a filled header band and row striping instead of
booktabs rules -- but it had been edited by hand and had drifted from the audited
source: a section's worth of results was missing, four calibrated parameters had
been dropped from an appendix, and a dose series had been quoted at a point that
was not its end.

This module keeps the better presentation and throws away the fork. The design
tokens below are read off that export; the content comes from the same Markdown
`neuromps.report.verify` checks, so the layout is a presentation choice and the
numbers stay auditable. `report/reference.docx` is that export with its body
emptied, and supplies the theme, the page geometry, the bullet numbering and the
footers.

The PDF is produced from the .docx by LibreOffice, so the two cannot disagree.

Usage:
    python -m neuromps.report.docx
    python -m neuromps.report.docx --source report/LAB_NOTEBOOK.md
"""

from __future__ import annotations

import argparse
import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.table import _Cell
from docx.text.paragraph import Paragraph as DocxParagraph

# --- design tokens, read off the reviewer's export ----------------------------

FONT = "Calibri"
CODE_FONT = "Consolas"

TITLE_PT, TITLE_RGB = 19.0, "1F4E79"
SUBTITLE_PT, SUBTITLE_RGB = 10.5, "555555"
ACT_PT, ACT_RGB = 13.5, "2F5597"
SECTION_PT, SECTION_RGB = 11.5, "1F4E79"
BODY_PT, BODY_RGB = 10.5, "222222"
LEAD_IN_RGB = "156082"
CAPTION_PT = 8.5
TABLE_PT = 10.0

HEADER_FILL = "1F4E79"
HEADER_RGB = "FFFFFF"
BAND_FILL = "F7F9FB"
PLAIN_FILL = "FFFFFF"

# A4 less the template's 1701-twip side margins.
TEXT_WIDTH_IN = 5.9
# Average advance per character at table size: the body font is proportional, the
# code font is not, and the gap is what decides whether an identifier wraps.
CHAR_IN = TABLE_PT / 72 * 0.50
CODE_CHAR_IN = TABLE_PT / 72 * 0.62
CELL_PADDING_IN = 0.14

# --- block model --------------------------------------------------------------


@dataclass(frozen=True)
class Heading:
    """A Markdown ATX heading. Level 1 is a title or an act, level 2 a section."""

    level: int
    text: str


@dataclass(frozen=True)
class Body:
    """A paragraph, a bullet, or a numbered item, already joined from its lines."""

    text: str
    kind: str = "paragraph"  # paragraph | bullet | ordered | quote | display
    marker: str = ""  # the glyph drawn in the hanging indent, for list items


@dataclass(frozen=True)
class Table:
    """A pipe table. The first row is a header unless every cell in it is empty."""

    rows: tuple[tuple[str, ...], ...]


@dataclass(frozen=True)
class Figure:
    """An image with its caption and the width the Markdown asks for."""

    path: str
    caption: str
    width_fraction: float


Block = Heading | Body | Table | Figure

# --- parsing ------------------------------------------------------------------

# Not a fenced-div fence, which also opens with a colon and must stay on its own line.
OPENS_WITH_PUNCTUATION = re.compile(r"^[ \t]*(?!:::)[,.;:)%]")
IMAGE = re.compile(r"^!\[(?P<caption>.*)\]\((?P<path>[^)]+)\)(?:\{width=(?P<width>\d+)%\})?\s*$")
# Bold, italic, code, and Pandoc's `~subscript~`, which takes no whitespace.
INLINE = re.compile(r"(\*\*.+?\*\*|\*.+?\*|`[^`]+`|~[^~\s]+~)", re.S)
TABLE_RULE = re.compile(r"^\|[\s:|-]*-[\s:|-]*\|\s*$")
NUMERIC_CELL = re.compile(r"^[\d\s.,%±+\-−–—()\[\]/:µ]*\d[\d\s.,%±+\-−–—()\[\]/:µ]*$")


def strip_citations(text: str) -> str:
    """Remove the `<!-- metrics.json: ... -->` pointers without scarring the prose.

    Pandoc drops these but keeps the space around them, which leaves "21,887
    records ," in the output and a bare "." opening a line. A comment that spans
    lines is removed with the lines it occupies; an inline one is removed with
    the space before it, and a line left holding nothing else goes with it.
    """

    def drop_if_multiline(match: re.Match[str]) -> str:
        # Only the comment itself decides. Testing the whole match tested the
        # newline this pattern also consumes, so a citation ending a line looked
        # multi-line, was removed with that newline, and welded the next line on:
        # "peaks at 0.48" + "at 100 uM" became "0.48at 100 uM".
        return "" if "\n" in match.group(1) else match.group(0)

    text = re.sub(r"[ \t]*(<!--.*?-->)[ \t]*\n?", drop_if_multiline, text, flags=re.S)
    kept: list[str] = []
    for original in text.splitlines(keepends=True):
        line = re.sub(r"[ \t]*<!--.*?-->", "", original)
        if original.strip() and not line.strip():
            continue
        if kept and OPENS_WITH_PUNCTUATION.match(line) and kept[-1].strip():
            kept[-1] = kept[-1].rstrip("\n") + line.lstrip()
        else:
            kept.append(line)
    return "".join(kept)


def _flush(buffer: list[str], kind: str, blocks: list[Block], marker: str = "") -> None:
    """Join a hard-wrapped block's lines and append it, if it holds anything.

    Runs of spaces are collapsed: a citation removed from the end of one line
    leaves one behind when the next line is joined on.
    """
    joined = re.sub(r" {2,}", " ", " ".join(part.strip() for part in buffer)).strip()
    if joined:
        blocks.append(Body(joined, kind, marker))
    buffer.clear()


def parse_blocks(text: str) -> list[Block]:
    """Turn the report Markdown into the block sequence the renderer walks.

    The source is hard-wrapped, so a paragraph, a bullet and a numbered item all
    continue onto following lines; only a blank line, a heading, a table row or
    an image ends one.
    """
    blocks: list[Block] = []
    buffer: list[str] = []
    kind, marker = "paragraph", ""
    rows: list[tuple[str, ...]] = []
    display = False  # inside a `::: equation` div

    def close_table() -> None:
        if rows:
            blocks.append(Table(tuple(rows)))
            rows.clear()

    for line in strip_citations(text).splitlines():
        stripped = line.strip()
        if stripped.startswith("|"):
            _flush(buffer, kind, blocks, marker)
            if not TABLE_RULE.match(stripped):
                rows.append(tuple(cell.strip() for cell in stripped.strip("|").split("|")))
            continue
        close_table()
        if not stripped or stripped == "---":
            _flush(buffer, kind, blocks, marker)
        elif match := re.match(r"^(#{1,6}) (.+)$", stripped):
            _flush(buffer, kind, blocks, marker)
            blocks.append(Heading(len(match.group(1)), match.group(2).strip()))
        elif stripped.startswith(":::"):
            # Pandoc's fenced div: `::: equation` opens a displayed line, `:::` closes it.
            _flush(buffer, kind, blocks, marker)
            display = stripped[3:].strip() == "equation"
            kind, marker = ("display" if display else "paragraph"), ""
        elif image := IMAGE.match(stripped):
            _flush(buffer, kind, blocks, marker)
            width = int(image.group("width") or 100) / 100
            blocks.append(Figure(image.group("path"), image.group("caption"), width))
        elif match := re.match(r"^(- |\d+\. |> )(.*)$", stripped):
            opener = match.group(1)
            if opener == "> " and kind == "quote" and buffer:
                buffer.append(match.group(2))  # the disclaimer runs to a second line
                continue
            _flush(buffer, kind, blocks, marker)
            kind = {"-": "bullet", ">": "quote"}.get(opener[0], "ordered")
            marker = "" if kind == "quote" else ("\u2022" if kind == "bullet" else opener.strip())
            buffer.append(match.group(2))
        else:
            if not buffer:
                kind, marker = ("display" if display else "paragraph"), ""
            buffer.append(stripped)
    _flush(buffer, kind, blocks, marker)
    close_table()
    return blocks


# --- run-level rendering -------------------------------------------------------


def _style_run(run, *, size: float, rgb: str, bold: bool = False, italic: bool = False) -> None:
    """Apply a design token to one run. Every run in the document goes through here."""
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(rgb)
    run.bold = bold
    run.italic = italic


def add_rich_text(
    paragraph: DocxParagraph,
    text: str,
    *,
    size: float = BODY_PT,
    rgb: str = BODY_RGB,
    lead_in: bool = True,
) -> None:
    """Write Markdown emphasis into a paragraph as runs.

    A paragraph that opens with a bold lead-in takes the accent colour the export
    uses for them, which is what makes its running text scannable.
    """
    first = True
    for piece in INLINE.split(text):
        if not piece:
            continue
        bold = piece.startswith("**")
        italic = not bold and piece.startswith("*")
        code = piece.startswith("`")
        run = paragraph.add_run(_plain(piece))
        colour = LEAD_IN_RGB if (bold and first and lead_in) else rgb
        _style_run(run, size=size, rgb=colour, bold=bold, italic=italic)
        if code:
            run.font.name = CODE_FONT
            run.font.size = Pt(size - 1)
        run.font.subscript = _is_subscript(piece) or None
        first = False


def _is_subscript(piece: str) -> bool:
    """Whether an `INLINE` piece is a `~subscript~` span."""
    return len(piece) > 2 and piece[0] == piece[-1] == "~"


def _plain(piece: str) -> str:
    """The text of an `INLINE` piece without its markup."""
    return piece[1:-1] if _is_subscript(piece) else piece.strip("*`")


def _shade(cell: _Cell, fill: str) -> None:
    """Fill a table cell. python-docx has no API for shading, so write the element."""
    shd = cell._tc.get_or_add_tcPr().makeelement(qn("w:shd"), {})
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


# Order is not decorative in `tblPr`: an element out of sequence makes the whole
# block invalid, and LibreOffice answers by ignoring it rather than complaining.
# Appending `tblBorders` after `tblLook` cost the tables their fixed widths once.
AFTER_BORDERS = ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook")


def _outer_border_only(table) -> None:
    """The export rules the table's outside edge and nothing within it."""
    borders = table._tbl.tblPr.makeelement(qn("w:tblBorders"), {})
    for edge in ("top", "left", "bottom", "right"):
        element = borders.makeelement(qn(f"w:{edge}"), {})
        element.set(qn("w:val"), "single")
        element.set(qn("w:color"), "auto")
        element.set(qn("w:sz"), "6")
        borders.append(element)
    table._tbl.tblPr.insert_element_before(borders, *AFTER_BORDERS)


def _keep_row_whole(row, with_next: bool) -> None:
    """Hold a row together, and to the row after it.

    `cantSplit` stops a striped row breaking mid-cell, where the two halves read
    as two rows. Keeping each row with the next stops a table starting at the
    foot of a page and leaving one row stranded there.
    """
    properties = row._tr.get_or_add_trPr()
    properties.append(properties.makeelement(qn("w:cantSplit"), {}))
    for cell in row.cells:
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.keep_with_next = with_next


def _measure(text: str) -> tuple[float, float]:
    """Return a cell's total width and its longest unbreakable word, in inches.

    Code spans are set in a monospace face whose advance is wider than the body
    font's average, so `pct_spikes_in_bursts` needs more room than its twenty
    characters suggest -- which is how it came to wrap mid-identifier.
    """
    total = longest = 0.0
    for piece in INLINE.split(text):
        if not piece:
            continue
        plain = _plain(piece)
        advance = CODE_CHAR_IN if piece.startswith("`") else CHAR_IN
        total += len(plain) * advance
        for word in plain.split():
            longest = max(longest, len(word) * advance)
    return total, longest


def column_shares(rows: tuple[tuple[str, ...], ...]) -> list[float]:
    """Split the text column between table columns, widest-word first.

    Word's autofit gives a column of four-digit numbers the same room as a column
    of sentences. Each column is first given enough to hold its longest single
    word unbroken, and only the width left over is shared out by how much text
    each column actually carries.
    """
    measured = [[_measure(row[column]) for row in rows] for column in range(len(rows[0]))]
    minimums = [max(word for _, word in column) + CELL_PADDING_IN for column in measured]
    demands = [max(total for total, _ in column) for column in measured]
    if sum(minimums) >= TEXT_WIDTH_IN:  # a table too wide to satisfy: share it out flat
        return [minimum / sum(minimums) for minimum in minimums]
    spare = TEXT_WIDTH_IN - sum(minimums)
    widths = [
        minimum + spare * demand / sum(demands)
        for minimum, demand in zip(minimums, demands, strict=True)
    ]
    return [width / TEXT_WIDTH_IN for width in widths]


def _apply_widths(table, shares: list[float]) -> None:
    """Fix the layout at those shares.

    All three of these are needed and none is redundant: a table left at
    `tblW auto` is sized from its content whatever the cells ask for, the grid is
    what a renderer lays out against, and Word reads the width from the cell
    rather than the column. Setting only the cells leaves the widths ignored.
    """
    table.autofit = False  # python-docx writes the `tblLayout` element in sequence
    width = table._tbl.tblPr.find(qn("w:tblW"))
    width.set(qn("w:type"), "dxa")
    width.set(qn("w:w"), str(round(TEXT_WIDTH_IN * 1440)))
    for column, share in enumerate(shares):
        inches = Inches(TEXT_WIDTH_IN * share)
        table.columns[column].width = inches
        for row in table.rows:
            row.cells[column].width = inches


def _repeat_header(row) -> None:
    """Carry the header band onto a continuation page, for the long appendix tables."""
    properties = row._tr.get_or_add_trPr()
    properties.append(properties.makeelement(qn("w:tblHeader"), {}))


# --- block-level rendering -----------------------------------------------------


def render_heading(document: Document, heading: Heading, first: bool) -> None:
    """Title, act or section, by the three heading looks the export uses."""
    if first:
        size, rgb, before, after = TITLE_PT, TITLE_RGB, 0, 4
    elif heading.level == 1:
        size, rgb, before, after = ACT_PT, ACT_RGB, 18, 6
    else:
        size, rgb, before, after = SECTION_PT, SECTION_RGB, 12, 3
    if heading.text.lower().startswith("appendices"):
        document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.keep_with_next = True
    add_rich_text(paragraph, heading.text, size=size, rgb=rgb, lead_in=False)
    for run in paragraph.runs:
        run.bold = True


def render_body(document: Document, body: Body, subtitle: bool) -> None:
    """A paragraph, a bullet, a numbered item, or the regulatory blockquote.

    List markers are drawn as a run in a hanging indent rather than through a
    numbered style: the template's `List Bullet` carries no numbering definition,
    so a style-only bullet renders as a bare indent in Word and in LibreOffice
    alike -- which is how the first build lost every bullet glyph silently.
    """
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    if body.marker:
        paragraph.paragraph_format.left_indent = Inches(0.3)
        paragraph.paragraph_format.first_line_indent = Inches(-0.18)
        _style_run(paragraph.add_run(f"{body.marker}\t"), size=BODY_PT, rgb=BODY_RGB)
    if body.kind == "display":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(2)
        add_rich_text(paragraph, body.text, size=BODY_PT + 1, lead_in=False)
        return
    if body.kind == "quote":
        paragraph.paragraph_format.left_indent = Inches(0.25)
        add_rich_text(paragraph, body.text, size=BODY_PT, rgb=SUBTITLE_RGB, lead_in=False)
        for run in paragraph.runs:
            run.italic = True
        return
    if subtitle:
        add_rich_text(paragraph, body.text, size=SUBTITLE_PT, rgb=SUBTITLE_RGB, lead_in=False)
        for run in paragraph.runs:
            run.italic = True
        return
    add_rich_text(paragraph, body.text)


def _is_centred(text: str) -> bool:
    """Numbers, dashes and empty cells are centred; prose is set left."""
    plain = re.sub(r"[*`]", "", text)
    return bool(NUMERIC_CELL.match(plain)) or plain in {"", "—"}


def _fill_cell(
    cell: _Cell, text: str, *, header: bool, band: bool, first_column: bool, centre: bool
) -> None:
    """One table cell, in the header band or a body row.

    A header follows its column (`centre`): a centred label over left-set prose
    reads as misaligned, which is how a reviewer described the team table.
    """
    _shade(cell, HEADER_FILL if header else (BAND_FILL if band else PLAIN_FILL))
    cell.vertical_alignment = 1  # WD_ALIGN_VERTICAL.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    if centre if header else _is_centred(text):
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rgb = HEADER_RGB if header else BODY_RGB
    add_rich_text(paragraph, text, size=TABLE_PT, rgb=rgb, lead_in=False)
    for run in paragraph.runs:
        if header or first_column:
            run.bold = True
        if header:
            run.font.color.rgb = RGBColor.from_string(HEADER_RGB)


def render_table(document: Document, block: Table) -> None:
    """A pipe table in the export's design: filled header band, striped body.

    A table whose header cells are all empty -- the front-matter harness table is
    written that way -- is rendered as body rows throughout, so it does not open
    with an empty dark bar.
    """
    rows = block.rows
    has_header = any(cell.strip() for cell in rows[0])
    if not has_header:
        rows = rows[1:]  # an all-empty header row is how a headerless table is written
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _outer_border_only(table)
    _apply_widths(table, column_shares(rows))
    body = rows[1:] if has_header else rows
    centred = [all(_is_centred(row[column]) for row in body) for column in range(len(rows[0]))]
    for index, source in enumerate(rows):
        header = has_header and index == 0
        band = bool((index - (1 if has_header else 0)) % 2)
        for column, text in enumerate(source):
            _fill_cell(
                table.rows[index].cells[column],
                text,
                header=header,
                band=band,
                first_column=column == 0 and not header,
                centre=centred[column],
            )
    for index, row in enumerate(table.rows):
        _keep_row_whole(row, with_next=index < len(table.rows) - 1)
    if has_header:
        _repeat_header(table.rows[0])
    document.add_paragraph().paragraph_format.space_after = Pt(4)


def render_figure(document: Document, figure: Figure, source: Path, number: int) -> None:
    """Embed the image and caption it, at the width the Markdown asks for."""
    resolved = (source.parent / figure.path).resolve()
    if not resolved.exists():
        raise FileNotFoundError(f"figure referenced by {source.name} is missing: {figure.path}")
    holder = document.add_paragraph()
    holder.alignment = WD_ALIGN_PARAGRAPH.CENTER
    holder.paragraph_format.space_before = Pt(6)
    holder.paragraph_format.space_after = Pt(2)
    # An image at the foot of a page with its caption at the top of the next
    # reads as two unrelated things; the first full build did exactly that.
    holder.paragraph_format.keep_with_next = True
    holder.add_run().add_picture(str(resolved), width=Inches(TEXT_WIDTH_IN * figure.width_fraction))
    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(10)
    label = caption.add_run(f"Figure {number}. ")
    _style_run(label, size=CAPTION_PT, rgb=BODY_RGB, bold=True, italic=True)
    add_rich_text(caption, figure.caption, size=CAPTION_PT, rgb=BODY_RGB, lead_in=False)
    for run in caption.runs[1:]:
        run.italic = True


# --- document assembly ---------------------------------------------------------


def contents_rows(blocks: list[Block]) -> tuple[tuple[str, ...], ...]:
    """Build the act-by-act contents table out of the headings themselves.

    The export carried a hand-written version of this table, which is how its
    section list came to disagree with its sections. Deriving it from the
    document means it cannot.
    """
    rows: list[tuple[str, str]] = [("Act", "Sections")]
    act, sections = "", []
    for block in blocks[1:]:  # blocks[0] is the title, which is not an act
        if not isinstance(block, Heading):
            continue
        if block.level == 1:
            if act:
                rows.append((act, "\n".join(sections)))
            act, sections = block.text, []
        elif act and block.level == 2:
            sections.append(block.text)
    if act:
        rows.append((act, "\n".join(sections)))
    return tuple(rows)


def _render_contents(document: Document, blocks: list[Block]) -> None:
    """The contents table, with each act's sections stacked in one cell."""
    rows = contents_rows(blocks)
    table = document.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _outer_border_only(table)
    _apply_widths(table, [0.36, 0.64])
    for index, (act, sections) in enumerate(rows):
        header = index == 0
        for column, text in enumerate((act, sections)):
            cell = table.rows[index].cells[column]
            _shade(cell, HEADER_FILL if header else (BAND_FILL if index % 2 == 0 else PLAIN_FILL))
            cell.vertical_alignment = 1
            for line_number, line in enumerate(text.split("\n")):
                paragraph = cell.paragraphs[0] if line_number == 0 else cell.add_paragraph()
                paragraph.paragraph_format.space_after = Pt(0)
                run = paragraph.add_run(line)
                _style_run(
                    run,
                    size=TABLE_PT,
                    rgb=HEADER_RGB if header else BODY_RGB,
                    bold=header or column == 0,
                )
    for index, row in enumerate(table.rows):
        _keep_row_whole(row, with_next=index < len(table.rows) - 1)
    _repeat_header(table.rows[0])
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _keep_with_what_it_introduces(document: Document, blocks: list[Block], index: int) -> None:
    """Hold a paragraph to the table or figure that follows it.

    Tables are kept whole, so a long one jumps to the next page and leaves the
    sentence introducing it - and the heading above that - stranded at the foot
    of the previous one. Two appendices came out that way.
    """
    following = blocks[index + 1] if index + 1 < len(blocks) else None
    if isinstance(following, (Table, Figure)):
        document.paragraphs[-1].paragraph_format.keep_with_next = True


def render(source: Path, template: Path, output: Path) -> Path:
    """Render one Markdown report into Word and return the path written.

    The subtitle look is given to the one paragraph that follows the title, which
    is where every report in `report/` states what it is.
    """
    blocks = parse_blocks(source.read_text(encoding="utf-8"))
    document = Document(str(template))
    seen_title = False
    seen_contents = False
    figure_number = 0
    for index, block in enumerate(blocks):
        if isinstance(block, Heading):
            if block.level == 1 and seen_title and not seen_contents:
                _render_contents(document, blocks)
                seen_contents = True
            render_heading(document, block, first=not seen_title)
            seen_title = True
        elif isinstance(block, Body):
            render_body(document, block, subtitle=index == 1)
            _keep_with_what_it_introduces(document, blocks, index)
        elif isinstance(block, Table):
            render_table(document, block)
        else:
            figure_number += 1
            render_figure(document, block, source, figure_number)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output))
    return output


def main(argv: list[str] | None = None) -> int:
    """Render a report to .docx. The PDF is produced from it by report/build_docx.sh."""
    root = Path(__file__).resolve().parents[3]
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=root / "report/TECHNICAL_REPORT.md")
    parser.add_argument("--template", type=Path, default=root / "report/reference.docx")
    parser.add_argument("--output", type=Path, default=None)
    arguments = parser.parse_args(argv)
    output = arguments.output or arguments.source.with_suffix(".docx")
    written = render(arguments.source, arguments.template, output)
    print(f"docx: wrote {written} ({written.stat().st_size // 1024} KB)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
