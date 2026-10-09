"""The Word build must carry the audited source, not a hand-edited fork of it.

The reason this renderer exists is that a Word export of the report was edited by
hand and drifted: a section lost its closing results, an appendix lost four of the
parameters it documented, and a table of contents went on listing sections that had
moved. These tests hold the rendered document to the Markdown it came from.
"""

from __future__ import annotations

import re
from pathlib import Path

import pytest

pytest.importorskip("docx")

from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

from neuromps.report.docx import (  # noqa: E402
    CELL_PADDING_IN,
    CODE_CHAR_IN,
    TEXT_WIDTH_IN,
    Body,
    Figure,
    Heading,
    Table,
    add_rich_text,
    column_shares,
    contents_rows,
    parse_blocks,
    render,
    strip_citations,
)

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "report/TECHNICAL_REPORT.md"
TEMPLATE = ROOT / "report/reference.docx"


def test_a_citation_is_removed_without_scarring_the_sentence() -> None:
    """The pointers are apparatus, not prose, and must not leave punctuation adrift."""
    text = "It scores 0.3482 <!-- metrics.json: cross_lab_transfer.balanced_accuracy -->, against\n"
    assert strip_citations(text).strip() == "It scores 0.3482, against"


def test_a_citation_on_its_own_line_takes_the_line_with_it() -> None:
    """A comment that begins a line would otherwise leave a blank splitting a paragraph."""
    text = (
        "of the nine mechanism classes\n"
        "<!-- metrics.json: pooling.n_classes -->\n"
        "appear in one source\n"
    )
    assert strip_citations(text) == "of the nine mechanism classes\nappear in one source\n"


def test_a_comment_spanning_lines_is_removed_whole() -> None:
    """The team-table note runs to four lines and must not survive as prose."""
    text = "**Team.**\n<!-- FILL IN\n     BEFORE\n     SUBMITTING -->\n\n| member |\n"
    assert "FILL IN" not in strip_citations(text)
    assert "**Team.**" in strip_citations(text)


def test_a_hard_wrapped_paragraph_becomes_one_block() -> None:
    """The source wraps at 100 columns; a paragraph is not three paragraphs."""
    blocks = parse_blocks("one two\nthree four\n\nnext\n")
    assert [b.text for b in blocks if isinstance(b, Body)] == ["one two three four", "next"]


def test_a_bullet_keeps_its_continuation_lines() -> None:
    """A wrapped bullet must not split into a bullet and a stray paragraph."""
    blocks = parse_blocks("- **Provenance.** The build fails if any\n  hash disagrees.\n")
    assert len(blocks) == 1
    assert isinstance(blocks[0], Body) and blocks[0].kind == "bullet"
    assert blocks[0].text == "**Provenance.** The build fails if any hash disagrees."


def test_a_table_drops_its_alignment_rule_and_keeps_its_rows() -> None:
    """The `|---|---|` separator is Markdown syntax, not a row of data."""
    blocks = parse_blocks("| a | b |\n|---|---|\n| 1 | 2 |\n")
    assert blocks == [Table((("a", "b"), ("1", "2")))]


def test_a_figure_carries_its_caption_and_requested_width() -> None:
    """Width is a fraction of the text column, so 62% must not render full-bleed."""
    blocks = parse_blocks("![A caption, REAL data.](../outputs/figures/x.png){width=62%}\n")
    assert blocks == [Figure("../outputs/figures/x.png", "A caption, REAL data.", 0.62)]


def test_the_contents_table_is_derived_from_the_headings() -> None:
    """A hand-written contents table is how the export's list went stale."""
    blocks = [
        Heading(1, "Title"),
        Heading(1, "ACT I"),
        Heading(2, "1. First"),
        Heading(2, "2. Second"),
        Heading(1, "ACT II"),
        Heading(2, "3. Third"),
    ]
    assert contents_rows(blocks) == (
        ("Act", "Sections"),
        ("ACT I", "1. First\n2. Second"),
        ("ACT II", "3. Third"),
    )


@pytest.fixture(scope="module")
def rendered(tmp_path_factory: pytest.TempPathFactory) -> Document:
    """Render the real report once; the checks below all read the same document."""
    output = tmp_path_factory.mktemp("docx") / "TECHNICAL_REPORT.docx"
    return Document(str(render(SOURCE, TEMPLATE, output)))


def test_every_figure_in_the_source_is_embedded(rendered: Document) -> None:
    """The export lost all of its figures silently; this is the check that catches it."""
    wanted = len(re.findall(r"^!\[", SOURCE.read_text(encoding="utf-8"), re.M))
    parts = rendered.part.package.parts
    embedded = sum(1 for part in parts if part.content_type.startswith("image/"))
    assert wanted > 0
    assert embedded == wanted


def test_no_citation_pointer_survives_into_the_document(rendered: Document) -> None:
    """A metrics pointer rendered as visible text would be a defect a reader sees."""
    text = "\n".join(p.text for p in rendered.paragraphs)
    assert "metrics.json:" not in text
    assert "<!--" not in text


def test_the_sections_present_in_the_source_are_present_in_the_document(
    rendered: Document,
) -> None:
    """The fork lost §5's closing results. Every heading must survive the render."""
    headings = [
        block.text
        for block in parse_blocks(SOURCE.read_text(encoding="utf-8"))
        if isinstance(block, Heading)
    ]
    rendered_text = "\n".join(p.text for p in rendered.paragraphs)
    missing = [h for h in headings if re.sub(r"[*`]", "", h) not in rendered_text]
    assert not missing, f"headings lost in the render: {missing}"


def test_every_source_table_row_reaches_the_document(rendered: Document) -> None:
    """An appendix lost four rows in the fork; row counts are checked, not eyeballed."""
    source = parse_blocks(SOURCE.read_text(encoding="utf-8"))
    source_tables = [b for b in source if isinstance(b, Table)]
    # One extra table is generated: the contents table derived from the headings.
    tables = rendered.tables  # python-docx rebuilds these objects per access, so bind once
    generated = [t for t in tables if t.rows[0].cells[0].text == "Act"]
    assert len(generated) == 1
    assert len(tables) == len(source_tables) + 1
    assert len(generated[0].rows) == len(contents_rows(source))
    # A table written with an empty header row drops it rather than opening on a
    # blank dark band, so that one row is expected to be absent.
    expected = [len(t.rows) - (0 if any(t.rows[0]) else 1) for t in source_tables]
    carried = [len(t.rows) for t in tables if t is not generated[0]]
    assert carried == expected


def test_the_table_header_band_is_filled_and_white(rendered: Document) -> None:
    """The header band is the reason this renderer exists; assert it is actually set."""
    header = rendered.tables[0].rows[0].cells[0]
    assert 'w:fill="1F4E79"' in header._tc.xml
    runs = [run for p in header.paragraphs for run in p.runs]
    assert runs and all(str(run.font.color.rgb) == "FFFFFF" and run.bold for run in runs)


def test_a_table_written_without_a_header_gets_no_dark_band(rendered: Document) -> None:
    """The front-matter harness table has an empty header row; it must not render one."""
    harness = next(
        t for t in rendered.tables if any(r.cells[0].text == "Verification harness" for r in t.rows)
    )
    assert 'w:fill="1F4E79"' not in harness.rows[0].cells[0]._tc.xml


def test_every_list_item_carries_a_visible_marker(rendered: Document) -> None:
    """The template's List Bullet style has no numbering, so the glyph is drawn."""
    wanted = sum(
        1
        for block in parse_blocks(SOURCE.read_text(encoding="utf-8"))
        if isinstance(block, Body) and block.marker
    )
    markers = ("\u2022\t", "1.\t", "2.\t", "3.\t")
    drawn = sum(1 for p in rendered.paragraphs if p.text.startswith(markers))
    assert wanted > 0
    assert drawn == wanted


def test_a_column_is_wide_enough_for_its_longest_identifier() -> None:
    """`pct_spikes_in_bursts` wrapped mid-word until columns were sized by word, not cell."""
    rows = (
        ("feature", "rank agreement", "practical consequence"),
        ("`pct_spikes_in_bursts`", "0.9703", "reproduces on both axes, which is worth saying"),
    )
    shares = column_shares(rows)
    assert sum(shares) == pytest.approx(1.0)
    identifier = len("pct_spikes_in_bursts") * CODE_CHAR_IN
    assert shares[0] * TEXT_WIDTH_IN >= identifier + CELL_PADDING_IN


def test_the_table_layout_is_actually_fixed(rendered: Document) -> None:
    """Order in `tblPr` is not decorative.

    Appending `tblBorders` after `tblLook` and a second `tblLayout` made the block
    invalid; LibreOffice answered by ignoring it and auto-fitting, so the computed
    widths were silently discarded and the tables looked untouched.
    """
    order = ["tblW", "jc", "tblBorders", "tblLayout", "tblLook"]
    for table in rendered.tables:
        tags = [child.tag.split("}")[1] for child in table._tbl.tblPr]
        assert tags == order, f"tblPr out of schema order: {tags}"
        layout = table._tbl.tblPr.find(qn("w:tblLayout"))
        assert layout.get(qn("w:type")) == "fixed"
        assert table._tbl.tblPr.find(qn("w:tblW")).get(qn("w:type")) == "dxa"
        assert all(cell.width is not None for cell in table.rows[0].cells)


def test_a_citation_ending_a_line_does_not_weld_the_next_line_on() -> None:
    """The stripper consumes the newline after a comment, so only the comment may decide.

    Testing the whole match tested that newline too: a citation at the end of a
    line looked like a multi-line comment, was removed along with the line break,
    and ran the following line into it as "0.48at 100 uM".
    """
    text = (
        "peaks at 0.48 <!-- metrics.json: a.b -->\n"
        "at 100 uM <!-- metrics.json: a.c -->\n"
        "and falls\n"
    )
    assert strip_citations(text) == "peaks at 0.48\nat 100 uM\nand falls\n"


def test_a_figure_image_is_kept_with_its_caption(rendered: Document) -> None:
    """An image at a page foot with its caption on the next page reads as two things."""
    holders = [p for p in rendered.paragraphs if p._p.xpath(".//pic:pic") and not p.text.strip()]
    assert holders, "no image paragraphs found"
    assert all(p.paragraph_format.keep_with_next for p in holders)


def test_the_sentence_introducing_a_table_stays_with_it(rendered: Document) -> None:
    """Tables are kept whole, so an unanchored intro gets stranded when one jumps a page."""
    intro = next(
        p for p in rendered.paragraphs if p.text.startswith("Calibrated values are the chosen draw")
    )
    assert intro.paragraph_format.keep_with_next


def test_a_tilde_span_is_set_as_a_true_subscript() -> None:
    """Pandoc's `x~sub~`. The Unicode glyph that stood in for it, a subscript e,
    printed so like a subscript zero that a reviewer read the definition as x0/x0."""
    document = Document()
    add_rich_text(document.add_paragraph(), "*x*~exposure~ / *x*~baseline~")
    runs = [(r.text, bool(r.font.subscript), bool(r.italic)) for r in document.paragraphs[0].runs]
    assert runs == [
        ("x", False, True),
        ("exposure", True, False),
        (" / ", False, False),
        ("x", False, True),
        ("baseline", True, False),
    ]


def test_a_tilde_with_spaces_around_it_stays_text() -> None:
    """Only a span with no whitespace is a subscript, as in Pandoc."""
    document = Document()
    add_rich_text(document.add_paragraph(), "roughly ~ 9 minutes ~ end to end")
    assert document.paragraphs[0].text == "roughly ~ 9 minutes ~ end to end"


def test_an_equation_div_becomes_one_displayed_line() -> None:
    """A fenced `::: equation` div; Pandoc's LaTeX path renders it as a plain paragraph."""
    blocks = parse_blocks("text\n\n::: equation\nΔ = log₂(a / b)\n:::\n\nafter\n")
    assert blocks == [Body("text"), Body("Δ = log₂(a / b)", "display"), Body("after")]


def test_the_report_equation_is_centred_with_real_subscripts(rendered: Document) -> None:
    """The one definition every model rests on must not be misreadable."""
    equation = next(p for p in rendered.paragraphs if p.text.startswith("Δ = log"))
    assert equation.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert [r.text for r in equation.runs if r.font.subscript] == ["2", "exposure", "baseline"]


def test_a_header_lines_up_with_the_column_beneath_it(rendered: Document) -> None:
    """Centred headers over left-set text read as misaligned; numbers stay centred."""
    team = next(t for t in rendered.tables if t.rows[0].cells[0].text == "member")
    assert team.rows[0].cells[1].paragraphs[0].alignment != WD_ALIGN_PARAGRAPH.CENTER
    models = next(t for t in rendered.tables if t.rows[0].cells[0].text == "model")
    assert models.rows[0].cells[1].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
